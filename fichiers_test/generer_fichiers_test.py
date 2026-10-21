#!/usr/bin/env python3
"""
generer_fichiers_test.py — Genere deux fichiers Excel de test pour ObjectifSMART-AI.

  1) objectifs_strategiques_BMCI.xlsx   -> import via RH > "Importer la strategie"
  2) objectifs_collaborateurs_BMCI.xlsx -> import via RH > "Importer les objectifs"

Les en-tetes correspondent exactement aux colonnes reconnues par
hr_import_service.py (insensible a la casse, ligne 1 = en-tetes, donnees des la ligne 2).

Melange volontaire d'objectifs BIEN SMART (KPI chiffre + echeance) et FLOUS
(sans mesure ni date) pour bien voir l'analyse IA distinguer les deux.

Lancer :  python generer_fichiers_test.py
"""

import os

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

HERE = os.path.dirname(os.path.abspath(__file__))

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(bold=True, color="FFFFFF")


def _style_header(ws, ncols):
    """Met en forme la ligne d'en-tete + fige le volet + auto-largeur simple."""
    for col in range(1, ncols + 1):
        cell = ws.cell(row=1, column=col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center")
    ws.freeze_panes = "A2"
    # largeurs lisibles
    widths = [10, 26, 30, 70, 22, 18, 14]
    from openpyxl.utils import get_column_letter

    for i, w in enumerate(widths[:ncols], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


# ============================================================================
# 1) OBJECTIFS STRATEGIQUES (catalogue entreprise -> utilise pour l'alignement)
#    Colonnes reconnues : ID | Titre | Description | Direction | Priorite
# ============================================================================
def build_strategie():
    wb = Workbook()
    ws = wb.active
    ws.title = "Strategie"
    ws.append(["ID", "Titre", "Description", "Direction", "Priorite"])

    rows = [
        [
            "STRAT-01",
            "Accelerer la transformation digitale",
            "Digitaliser 80% des parcours clients et atteindre 1 million d'utilisateurs actifs de l'app mobile d'ici fin 2026.",
            "Direction Digitale",
            1,
        ],
        [
            "STRAT-02",
            "Ameliorer la satisfaction client",
            "Atteindre un NPS de 60 et reduire le taux de reclamation de 25% avant decembre 2026.",
            "Direction Clientele",
            1,
        ],
        [
            "STRAT-03",
            "Renforcer la maitrise des risques",
            "Maintenir le cout du risque sous 0,5% des encours et reduire les creances en souffrance de 15% en 2026.",
            "Direction des Risques",
            1,
        ],
        [
            "STRAT-04",
            "Developper le financement des PME",
            "Augmenter de 15% l'encours de credit aux PME et financer 500 nouvelles entreprises en 2026.",
            "Direction Entreprises",
            2,
        ],
        [
            "STRAT-05",
            "Promouvoir la finance durable (ESG)",
            "Financer 2 milliards MAD de projets verts d'ici 2027 et publier un rapport ESG annuel.",
            "Direction RSE",
            2,
        ],
        [
            "STRAT-06",
            "Optimiser l'efficacite operationnelle",
            "Reduire le coefficient d'exploitation a 50% d'ici fin 2026 par l'automatisation des processus.",
            "Direction Financiere",
            2,
        ],
        [
            "STRAT-07",
            "Renforcer la cybersecurite",
            "Atteindre 100% de conformite aux normes Bank Al-Maghrib et reduire les incidents de securite de 40% en 2026.",
            "DSI",
            1,
        ],
        [
            "STRAT-08",
            "Developper les talents et l'engagement",
            "Former 100% des collaborateurs aux competences digitales et porter le taux d'engagement a 80% d'ici fin 2026.",
            "DRH",
            3,
        ],
    ]
    for r in rows:
        ws.append(r)

    _style_header(ws, 5)
    out = os.path.join(HERE, "objectifs_strategiques_BMCI.xlsx")
    wb.save(out)
    return out, len(rows)


# ============================================================================
# 2) OBJECTIFS COLLABORATEURS
#    Colonnes : ID | Collaborateur | Email | Titre | Description | Entite | Manager | Periode | Contexte
#    -> les comptes COLLAB manquants sont crees a la volee (mdp temporaire fort).
# ============================================================================
def build_collaborateurs():
    wb = Workbook()
    ws = wb.active
    ws.title = "Objectifs"
    ws.append(
        [
            "ID",
            "Collaborateur",
            "Email",
            "Titre",
            "Description",
            "Entite",
            "Manager",
            "Periode",
            "Contexte",
        ]
    )

    # SMART = objectif clair, mesurable (chiffre) et date -> doit bien scorer
    # FLOU  = sans KPI ni echeance -> l'IA doit le signaler
    rows = [
        # --- Yassine (Reseau Agences) ---
        [
            "OBJ-001",
            "Yassine El Amrani",
            "yassine.elamrani@bmci.ma",
            "Augmenter le taux d'equipement digital des clients de l'agence de 45% a 65% d'ici le 31 decembre 2026",
            "Promouvoir l'app mobile BMCI aupres des clients particuliers lors de chaque rendez-vous.",
            "Reseau Agences",
            "Khalid Bennani",
            "2026",
            "Agence Casablanca Maarif",
        ],  # SMART
        [
            "OBJ-002",
            "Yassine El Amrani",
            "yassine.elamrani@bmci.ma",
            "Ameliorer la relation client",
            "Etre plus a l'ecoute et plus disponible pour les clients.",
            "Reseau Agences",
            "Khalid Bennani",
            "2026",
            "Agence Casablanca Maarif",
        ],  # FLOU
        # --- Fatima Zahra (Entreprises) ---
        [
            "OBJ-003",
            "Fatima Zahra Idrissi",
            "fatima.idrissi@bmci.ma",
            "Reduire le delai moyen de traitement des dossiers de credit de 8 a 5 jours ouvres avant fin du T3 2026",
            "Optimiser le circuit d'instruction et de validation des credits entreprises.",
            "Direction Entreprises",
            "Nadia Tazi",
            "T3 2026",
            "Pole credits PME",
        ],  # SMART
        [
            "OBJ-004",
            "Fatima Zahra Idrissi",
            "fatima.idrissi@bmci.ma",
            "Developper le portefeuille PME",
            "Prospecter de nouvelles entreprises et renforcer les relations existantes.",
            "Direction Entreprises",
            "Nadia Tazi",
            "2026",
            "Pole credits PME",
        ],  # FLOU
        # --- Omar (Reseau Agences) ---
        [
            "OBJ-005",
            "Omar Benjelloun",
            "omar.benjelloun@bmci.ma",
            "Ouvrir 50 nouveaux comptes professionnels par trimestre tout au long de 2026",
            "Conquete de la clientele professionnelle et artisans.",
            "Reseau Agences",
            "Khalid Bennani",
            "2026",
            "Agence Rabat Agdal",
        ],  # SMART
        [
            "OBJ-006",
            "Omar Benjelloun",
            "omar.benjelloun@bmci.ma",
            "Monter en competence sur les nouveaux produits",
            "Mieux connaitre l'offre pour mieux conseiller.",
            "Reseau Agences",
            "Khalid Bennani",
            "T1 2026",
            "Agence Rabat Agdal",
        ],  # FLOU
        # --- Salma (Conformite) ---
        [
            "OBJ-007",
            "Salma Cherkaoui",
            "salma.cherkaoui@bmci.ma",
            "Porter le score de conformite de l'equipe de 75% a 95% et obtenir la certification AMF d'ici juin 2026",
            "Renforcer la conformite reglementaire et la lutte anti-blanchiment.",
            "Direction Conformite",
            "Rachid Alaoui",
            "S1 2026",
            "Pole LCB-FT",
        ],  # SMART
        [
            "OBJ-008",
            "Salma Cherkaoui",
            "salma.cherkaoui@bmci.ma",
            "Ameliorer la qualite des controles",
            "Realiser des controles plus reguliers et mieux documentes.",
            "Direction Conformite",
            "Rachid Alaoui",
            "2026",
            "Pole LCB-FT",
        ],  # FLOU
        # --- Mehdi (DSI) ---
        [
            "OBJ-009",
            "Mehdi Tahiri",
            "mehdi.tahiri@bmci.ma",
            "Migrer 100% des serveurs critiques vers le nouveau datacenter securise avant le 30 septembre 2026",
            "Projet d'infrastructure pour la resilience et la cybersecurite.",
            "DSI",
            "Samir Ouazzani",
            "T3 2026",
            "Programme Resilience IT",
        ],  # SMART
        [
            "OBJ-010",
            "Mehdi Tahiri",
            "mehdi.tahiri@bmci.ma",
            "Reduire les incidents de securite",
            "Ameliorer la surveillance et la reactivite face aux menaces.",
            "DSI",
            "Samir Ouazzani",
            "2026",
            "Programme Resilience IT",
        ],  # FLOU
    ]
    for r in rows:
        ws.append(r)

    _style_header(ws, 9)
    out = os.path.join(HERE, "objectifs_collaborateurs_BMCI.xlsx")
    wb.save(out)
    return out, len(rows)


if __name__ == "__main__":
    p1, n1 = build_strategie()
    p2, n2 = build_collaborateurs()
    print(f"[OK] {n1} objectifs strategiques  -> {p1}")
    print(f"[OK] {n2} objectifs collaborateurs -> {p2}")
    print(
        "\nDans l'app : connecte-toi en RH (compte seede, voir backend/.env.local), onglet Imports :"
    )
    print("  - 'Importer la strategie'   -> objectifs_strategiques_BMCI.xlsx")
    print("  - 'Importer les objectifs'  -> objectifs_collaborateurs_BMCI.xlsx")
