#!/usr/bin/env python3
"""
services/document_extractor.py — Extraction de texte des fichiers téléversés.

RÔLE
    Convertit le contenu binaire d'un fichier (PDF, Word, Excel, texte) en texte
    brut, prêt à être découpé puis indexé par `rag_service`. Le type est déduit de
    l'extension du nom de fichier. Les bibliothèques lourdes sont importées à la
    demande (lazy) pour ne charger que ce qui est nécessaire.

EMPLACEMENT MVC
    Couche "services" : utilitaire métier, sans route HTTP.
"""

import io
import logging
import re

logger = logging.getLogger(__name__)

# Extensions supportées (déduites du nom de fichier).
SUPPORTED_EXTENSIONS = (".pdf", ".docx", ".xlsx", ".txt", ".md")

# Détection d'extraction "pauvre" : un run de >=3 lettres = un mot ; un run de
# >=2 lettres = la preuve qu'une ligne porte du texte (et n'est pas qu'une puce).
_WORD_RE = re.compile(r"[^\W\d_]{3,}", re.UNICODE)
_HAS_WORD_RE = re.compile(r"[^\W\d_]{2,}", re.UNICODE)


def file_extension(filename: str) -> str:
    """Retourne l'extension en minuscules (point inclus), ou '' si absente."""
    name = (filename or "").lower().strip()
    dot = name.rfind(".")
    return name[dot:] if dot != -1 else ""


def extract_text(filename: str, content: bytes) -> str:
    """Extrait le texte d'un fichier selon son extension.

    Lève ValueError si le format n'est pas supporté ou si rien n'est extractible.
    """
    ext = file_extension(filename)
    if ext == ".pdf":
        text = _extract_pdf(content)
    elif ext == ".docx":
        text = _extract_docx(content)
    elif ext == ".xlsx":
        text = _extract_xlsx(content)
    elif ext in (".txt", ".md"):
        text = _extract_plain(content)
    else:
        raise ValueError(
            f"Format non supporté : '{ext or '?'}'. "
            f"Formats acceptés : {', '.join(SUPPORTED_EXTENSIONS)}."
        )

    text = (text or "").strip()
    if not text:
        raise ValueError("Aucun texte extractible de ce fichier.")
    if _looks_poorly_extracted(text):
        raise ValueError(
            "Le document a été lu mais son contenu semble illisible ou mal "
            "structuré (trop peu de texte exploitable — PDF scanné/en image ou "
            "mise en page complexe). Fournissez un PDF contenant du vrai texte, "
            "ou une version .docx / .txt du document."
        )
    return text


def _looks_poorly_extracted(text: str) -> bool:
    """Détecte une extraction appauvrie : document lu mais vidé de son contenu.

    Cas typique : un PDF à mise en page riche (guide infographique) dont
    l'extracteur ne récupère que les titres et une forêt de puces/coches sans
    le texte associé. Indexer un tel contenu donne au RAG un contexte inutile
    (le Coach répond alors de façon générique). On préfère refuser l'import avec
    un message clair plutôt que d'indexer silencieusement du vide.
    """
    lines = [ln.strip() for ln in text.splitlines() if ln.strip()]
    if not lines:
        return True
    # Structure effondrée : la majorité des lignes ne portent aucun mot
    # (puces « • », coches « ✓ », séparateurs, caractères isolés). C'est la
    # signature d'un PDF à mise en page riche dont seul le squelette a été lu.
    # On ne se fie PAS à un plancher de mots absolu : une note courte mais
    # valide (« Réunion lundi, préparer le budget ») serait rejetée à tort.
    symbol_lines = sum(1 for ln in lines if not _HAS_WORD_RE.search(ln))
    if len(lines) >= 8 and symbol_lines / len(lines) >= 0.5:
        return True
    # Contenu non négligeable mais quasi aucun mot réel (ligatures/garbage).
    if len("".join(text.split())) > 40 and len(_WORD_RE.findall(text)) < 5:
        return True
    return False


def _extract_plain(content: bytes) -> str:
    for encoding in ("utf-8", "latin-1"):
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue
    return content.decode("utf-8", errors="ignore")


def _extract_pdf(content: bytes) -> str:
    # PyMuPDF (fitz) : extraction bien plus robuste que pypdf sur les mises en
    # page complexes (colonnes, cadres, guides infographiques).
    try:
        import pymupdf  # PyMuPDF >= 1.24
    except ImportError:  # pragma: no cover - compat anciennes versions
        import fitz as pymupdf

    with pymupdf.open(stream=content, filetype="pdf") as doc:
        text = "\n\n".join((page.get_text("text") or "") for page in doc)
        # PDF en images (scan, guide infographique) : la couche texte est quasi
        # vide. On bascule sur l'OCR local (rapidocr) en rendant chaque page.
        if _looks_poorly_extracted(text):
            ocr_text = _ocr_pdf_pages(doc)
            if len(ocr_text.strip()) > len(text.strip()):
                return ocr_text
    return text


# Moteur OCR chargé paresseusement puis réutilisé (le chargement des modèles
# ONNX coûte ~1-2 s ; inutile de le refaire à chaque téléversement).
_ocr_engine = None
_OCR_DPI = 200  # compromis qualité/vitesse pour l'OCR des pages
_OCR_MAX_PAGES = 50  # borne anti-emballement sur de très gros documents


def _get_ocr_engine():
    global _ocr_engine
    if _ocr_engine is None:
        from rapidocr_onnxruntime import RapidOCR  # import tardif : lourd (onnx)

        _ocr_engine = RapidOCR()
    return _ocr_engine


def _ocr_pdf_pages(doc) -> str:
    """OCR local des pages d'un PDF (texte dans des images).

    Local et hors-ligne (CNDP) : aucune donnée ne sort. Si rapidocr n'est pas
    installé, on renvoie "" et l'appelant conserve le texte natif (le garde-fou
    d'extract_text rejettera alors proprement l'import avec un message clair).
    """
    try:
        engine = _get_ocr_engine()
    except ImportError:
        logger.warning("OCR indisponible : rapidocr-onnxruntime n'est pas installé.")
        return ""

    n_pages = min(doc.page_count, _OCR_MAX_PAGES)
    if doc.page_count > _OCR_MAX_PAGES:
        logger.warning(
            "OCR limité aux %d premières pages (document de %d pages).",
            _OCR_MAX_PAGES,
            doc.page_count,
        )
    parts = []
    for i in range(n_pages):
        try:
            pixmap = doc[i].get_pixmap(dpi=_OCR_DPI)
            result, _ = engine(pixmap.tobytes("png"))
        except Exception as exc:  # une page illisible ne doit pas tout annuler
            logger.warning("OCR page %d échouée : %s", i + 1, exc)
            continue
        if result:
            parts.append("\n".join(line[1] for line in result))
    text = "\n\n".join(parts)
    logger.info("OCR terminé : %d page(s), %d caractères extraits.", n_pages, len(text))
    return text


def _extract_docx(content: bytes) -> str:
    import docx

    document = docx.Document(io.BytesIO(content))
    parts = [p.text for p in document.paragraphs if p.text and p.text.strip()]
    # Tables Word : on aplatit chaque ligne (cellules séparées par « | »).
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text and c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n\n".join(parts)


def _extract_xlsx(content: bytes) -> str:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    parts = []
    for ws in wb.worksheets:
        parts.append(f"# Feuille : {ws.title}")
        for row in ws.iter_rows(values_only=True):
            cells = [str(c).strip() for c in row if c is not None and str(c).strip()]
            if cells:
                parts.append(" | ".join(cells))
    wb.close()
    return "\n".join(parts)
