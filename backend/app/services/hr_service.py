#!/usr/bin/env python3
"""
services/hr_service.py — Logique métier du domaine RH.

RÔLE
    Analyses personnelles RH, imports Excel (objectifs collaborateurs & objectifs
    stratégiques), analyse en lot (SSE), tableaux de bord analytiques, liste
    détaillée et exports Excel/PDF. ISOLATION STRICTE : le RH ne voit QUE les
    données de périmètre HR (data_scope='HR' ou import_batch_id LIKE 'obj-%').

    Toute la persistance passe par les repositories ; la génération de documents
    est déléguée à export_service.

EMPLACEMENT MVC
    Couche "services" : logique métier.
"""

import io
import json
import secrets
from datetime import datetime
from typing import List

from fastapi import HTTPException

from app.core.security import hash_password
from app.database.session import db_session
from app.repositories import (
    user_repository,
    objective_repository,
    company_objective_repository,
    analysis_repository,
    hr_import_repository,
)
from app.services.ai_service import analyze_smart_internal, analyze_alignment_internal
from app.services import export_service

try:
    import openpyxl
    EXCEL_SUPPORT = True
except ImportError:  # pragma: no cover
    EXCEL_SUPPORT = False


def _require_hr(current_user: dict) -> None:
    if current_user["role"] != "HR":
        raise HTTPException(status_code=403, detail="Accès réservé aux RH.")


def _smart_flags_score(smart_result: dict) -> float:
    return (
        (1 if smart_result.get("S") else 0)
        + (1 if smart_result.get("M") else 0)
        + (1 if smart_result.get("A") else 0)
        + (1 if smart_result.get("R") else 0)
        + (1 if smart_result.get("T") else 0)
    ) / 5 * 100


def _smart_values(smart_result: dict) -> dict:
    issues = smart_result.get("issues", [])
    return {
        "s_score": 1 if smart_result.get("S") else 0,
        "m_score": 1 if smart_result.get("M") else 0,
        "a_score": 1 if smart_result.get("A") else 0,
        "r_score": 1 if smart_result.get("R") else 0,
        "t_score": 1 if smart_result.get("T") else 0,
        "issues_count": len(issues),
        "top_issue": issues[0].get("problem", "") if issues else "",
    }


# ============= ANALYSE PERSONNELLE RH =============

def analyze_personal(request, current_user: dict) -> dict:
    """Analyse SMART d'un objectif personnel RH, stockée dans hr_analysis_results."""
    _require_hr(current_user)
    smart_result = analyze_smart_internal(request.title, request.description or "")

    with db_session() as db:
        company_objectives = company_objective_repository.list_for_alignment(db)

    if company_objectives:
        alignment_result = analyze_alignment_internal(request.title, company_objectives)
    else:
        alignment_result = {"alignment_score": 0, "top_matches": [], "gaps": ["Aucun objectif stratégique défini"]}

    smart_score = _smart_flags_score(smart_result)
    sv = _smart_values(smart_result)
    result_id = f"hr-personal-{secrets.token_hex(6)}"

    with db_session() as db:
        analysis_repository.insert_hr(db, {
            "id": result_id,
            "objective_id": f"hr-obj-{result_id}",
            "objective_title": request.title,
            "smart_analysis": json.dumps(smart_result),
            "alignment_analysis": json.dumps(alignment_result),
            "smart_score": smart_score,
            "alignment_score": alignment_result.get("alignment_score", 0),
            "analyzed_at": datetime.now().isoformat(),
            **sv,
        })

    return {
        "success": True,
        "result_id": result_id,
        "title": request.title,
        "description": request.description,
        "smart_analysis": smart_result,
        "alignment_analysis": alignment_result,
        "smart_score": smart_score,
        "scope": "HR_PERSONAL",
    }


def get_analysis_results(current_user: dict) -> dict:
    """Résultats d'analyse personnels du RH."""
    _require_hr(current_user)
    with db_session() as db:
        rows = analysis_repository.list_hr_personal_results(db)
        stats = analysis_repository.hr_personal_stats(db)

    results = []
    for result in rows:
        for key in ("smart_analysis", "alignment_analysis"):
            if result.get(key):
                try:
                    result[key] = json.loads(result[key])
                except (json.JSONDecodeError, TypeError):
                    result[key] = None
        if not result.get("objective_title"):
            result["objective_title"] = result.get("obj_title") or f"Objectif {result.get('objective_id', 'N/A')}"
        results.append(result)

    return {
        "results": results,
        "stats": {
            "total_analyzed": stats.get("total_analyzed", 0),
            "avg_smart_score": round(stats.get("avg_smart_score") or 0, 1),
            "avg_alignment_score": round(stats.get("avg_alignment_score") or 0, 1),
        },
        "scope": "HR_PERSONAL",
    }


# ============= IMPORTS EXCEL =============

def import_strategy(filename: str, content: bytes, current_user: dict) -> dict:
    """Import Excel des objectifs stratégiques d'entreprise."""
    _require_hr(current_user)
    if not EXCEL_SUPPORT:
        raise HTTPException(status_code=500, detail="openpyxl non installé. Exécutez: pip install openpyxl")

    try:
        wb = openpyxl.load_workbook(io.BytesIO(content))
        ws = wb.active
        headers = [str(cell.value or "").lower().strip() for cell in ws[1]]

        col_map = {}
        for idx, h in enumerate(headers):
            if h in ["id", "reference", "ref", "code"]:
                col_map["id"] = idx
            elif h in ["titre", "title", "objectif", "axe", "axe stratégique", "axe strategique"]:
                col_map["title"] = idx
            elif h in ["description", "desc", "detail", "détail"]:
                col_map["description"] = idx
            elif h in ["departement", "department", "direction", "entité", "entite", "entity"]:
                col_map["department"] = idx
            elif h in ["priorité", "priorite", "priority", "poids", "weight"]:
                col_map["priority"] = idx

        if "title" not in col_map:
            raise HTTPException(status_code=400, detail="Colonne 'titre' ou 'title' obligatoire dans le fichier Excel.")

        batch_id = f"strat-{secrets.token_hex(4)}"
        imported = 0

        with db_session() as db:
            for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
                title = row[col_map["title"]] if col_map.get("title") is not None else None
                if not title or str(title).strip() == "":
                    continue

                obj_id = row[col_map["id"]] if col_map.get("id") is not None else f"STRAT-{row_idx}"
                description = row[col_map["description"]] if col_map.get("description") is not None else ""
                department = row[col_map["department"]] if col_map.get("department") is not None else ""
                priority = row[col_map["priority"]] if col_map.get("priority") is not None else 1
                try:
                    priority = int(priority) if priority else 1
                except (ValueError, TypeError):
                    priority = 1

                company_objective_repository.merge(db, {
                    "id": str(obj_id).strip(),
                    "title": str(title).strip(),
                    "description": str(description or "").strip(),
                    "department": str(department or "").strip(),
                    "priority": priority,
                    "created_by": current_user["id"],
                    "import_batch_id": batch_id,
                })
                imported += 1

            hr_import_repository.create(
                db, import_id=batch_id, import_type="strategy",
                filename=filename, records_count=imported, user_id=current_user["id"],
            )

        return {"success": True, "batch_id": batch_id, "imported": imported, "filename": filename}

    except HTTPException:
        raise
    except Exception as e:
        print(f"[HR Import Error] {e}")
        raise HTTPException(status_code=400, detail=f"Erreur lors de l'import: {str(e)}")


def import_objectives(filename: str, content: bytes, current_user: dict) -> dict:
    """
    Import Excel des objectifs collaborateurs.

    SÉCURITÉ (F10) : les comptes créés au passage reçoivent un mot de passe
    aléatoire FORT par utilisateur (token_urlsafe), jamais un seed partagé. Les
    credentials temporaires sont renvoyés dans `created_users` pour transmission
    individuelle (à changer à la première connexion).
    """
    _require_hr(current_user)
    if not EXCEL_SUPPORT:
        raise HTTPException(status_code=500, detail="openpyxl non installé. Exécutez: pip install openpyxl")

    try:
        wb = openpyxl.load_workbook(io.BytesIO(content))
        ws = wb.active
        headers = [str(cell.value or "").lower().strip() for cell in ws[1]]

        col_map = {}
        for idx, h in enumerate(headers):
            if h in ["id", "reference", "ref"]:
                col_map["id"] = idx
            elif h in ["collaborateur", "collab", "user", "utilisateur", "nom", "name", "employee"]:
                col_map["user_name"] = idx
            elif h in ["email", "mail", "e-mail"]:
                col_map["email"] = idx
            elif h in ["titre", "title", "objectif", "objective"]:
                col_map["title"] = idx
            elif h in ["description", "desc", "detail", "détail"]:
                col_map["description"] = idx
            elif h in ["entité", "entite", "entity", "direction", "department", "service"]:
                col_map["entity"] = idx
            elif h in ["manager", "responsable", "n+1", "superviseur"]:
                col_map["manager"] = idx
            elif h in ["période", "periode", "period", "année", "annee", "year", "trimestre", "quarter"]:
                col_map["period"] = idx
            elif h in ["contexte", "context"]:
                col_map["context"] = idx

        if "title" not in col_map:
            raise HTTPException(status_code=400, detail="Colonne 'titre' ou 'title' obligatoire dans le fichier Excel.")

        batch_id = f"obj-{secrets.token_hex(4)}"
        imported = 0
        users_created = 0
        created_users: List[dict] = []

        with db_session() as db:
            for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
                title = row[col_map["title"]] if col_map.get("title") is not None else None
                if not title or str(title).strip() == "":
                    continue

                user_name = row[col_map["user_name"]] if col_map.get("user_name") is not None else f"Collaborateur-{row_idx}"
                email = row[col_map["email"]] if col_map.get("email") is not None else None
                entity = row[col_map["entity"]] if col_map.get("entity") is not None else ""
                manager = row[col_map["manager"]] if col_map.get("manager") is not None else ""
                period = row[col_map["period"]] if col_map.get("period") is not None else ""
                description = row[col_map["description"]] if col_map.get("description") is not None else ""
                context = row[col_map["context"]] if col_map.get("context") is not None else ""
                obj_id_raw = row[col_map["id"]] if col_map.get("id") is not None else None

                if email:
                    existing = user_repository.get_by_email(db, str(email).strip())
                    if existing:
                        user_id = existing.id
                    else:
                        user_id = f"user-{secrets.token_hex(4)}"
                        temp_password = secrets.token_urlsafe(16)
                        user_repository.create(
                            db, user_id=user_id, email=str(email).strip(),
                            password_hash=hash_password(temp_password),
                            name=str(user_name).strip(), role="COLLAB",
                            entity=str(entity or "").strip(),
                        )
                        users_created += 1
                        created_users.append({
                            "user_id": user_id,
                            "name": str(user_name).strip(),
                            "email": str(email).strip(),
                            "temporary_password": temp_password,
                        })
                else:
                    user_id = f"import-user-{row_idx}"
                    if not user_repository.get_by_id(db, user_id):
                        temp_password = secrets.token_urlsafe(16)
                        temp_email = f"import_{row_idx}_{batch_id}@temp.local"
                        user_repository.create(
                            db, user_id=user_id, email=temp_email,
                            password_hash=hash_password(temp_password),
                            name=str(user_name).strip(), role="COLLAB",
                            entity=str(entity or "").strip(),
                        )
                        users_created += 1
                        created_users.append({
                            "user_id": user_id,
                            "name": str(user_name).strip(),
                            "email": temp_email,
                            "temporary_password": temp_password,
                            "note": "Email synthetique : a editer avant transmission.",
                        })

                obj_id = str(obj_id_raw).strip() if obj_id_raw else f"obj-{secrets.token_hex(6)}"
                objective_repository.merge(db, {
                    "id": obj_id,
                    "user_id": user_id,
                    "title": str(title).strip(),
                    "description": str(description or "").strip(),
                    "context": str(context or "").strip(),
                    "entity": str(entity or "").strip(),
                    "manager_name": str(manager or "").strip(),
                    "period": str(period or "").strip(),
                    "import_batch_id": batch_id,
                    "data_scope": "HR",
                    "status": "draft",
                    "progress": "not_started",
                })
                imported += 1

            hr_import_repository.create(
                db, import_id=batch_id, import_type="objectives",
                filename=filename, records_count=imported, user_id=current_user["id"],
            )

        return {
            "success": True,
            "batch_id": batch_id,
            "imported": imported,
            "users_created": users_created,
            "filename": filename,
            "created_users": created_users,
        }

    except HTTPException:
        raise
    except Exception as e:
        print(f"[HR Import Error] {e}")
        raise HTTPException(status_code=400, detail=f"Erreur lors de l'import: {str(e)}")


# ============= ANALYSE EN LOT (SSE) =============

def build_batch_generator(objective_ids: List[str], current_user: dict):
    """
    Prépare l'analyse en lot et retourne un générateur SSE (text/event-stream).
    Seuls les objectifs de périmètre HR sont analysés.
    """
    _require_hr(current_user)

    with db_session() as db:
        if not objective_ids:
            objs = objective_repository.list_unanalyzed_hr(db, limit=100)
        else:
            objs = objective_repository.list_hr_by_ids(db, objective_ids)
        objectives = [{"id": o.id, "title": o.title, "description": o.description} for o in objs]
        company_objectives = company_objective_repository.list_for_alignment(db)

    total = len(objectives)
    user_id = current_user["id"]

    def generate():
        results = []
        analyzed = 0
        errors = 0
        yield f"data: {json.dumps({'type': 'start', 'total': total})}\n\n"

        for idx, obj in enumerate(objectives):
            yield f"data: {json.dumps({'type': 'progress', 'current': idx + 1, 'total': total, 'title': obj['title'][:80]})}\n\n"
            try:
                try:
                    smart_result = analyze_smart_internal(obj["title"], obj.get("description", ""))
                except Exception:
                    smart_result = {"S": False, "M": False, "A": False, "R": False, "T": False, "overall": False, "confidence": 0, "issues": []}

                if company_objectives:
                    try:
                        alignment_result = analyze_alignment_internal(obj["title"], company_objectives)
                    except Exception:
                        alignment_result = {"alignment_score": 0, "top_matches": [], "gaps": []}
                else:
                    alignment_result = {"alignment_score": 0, "top_matches": [], "gaps": ["Aucun objectif stratégique défini"]}

                smart_score = _smart_flags_score(smart_result)
                sv = _smart_values(smart_result)

                with db_session() as db:
                    objective_repository.update_fields(db, obj["id"], {
                        "smart_analysis": json.dumps(smart_result),
                        "alignment_analysis": json.dumps(alignment_result),
                        "analyzed_by_role": "HR",
                        "analyzed_by_user_id": user_id,
                        "data_scope": "HR",
                        "status": "analyzed",
                        "updated_at": datetime.now().isoformat(),
                    })
                    analysis_repository.upsert_hr(db, {
                        "id": f"res-{obj['id']}",
                        "objective_id": obj["id"],
                        "smart_score": smart_score,
                        "alignment_score": alignment_result.get("alignment_score", 0),
                        "issues_count": sv["issues_count"],
                        "top_issue": sv["top_issue"],
                        "analyzed_at": datetime.now().isoformat(),
                        "s_score": sv["s_score"], "m_score": sv["m_score"], "a_score": sv["a_score"],
                        "r_score": sv["r_score"], "t_score": sv["t_score"],
                    })

                results.append({
                    "id": obj["id"], "title": obj["title"],
                    "smart_score": smart_score,
                    "alignment_score": alignment_result.get("alignment_score", 0),
                    "success": True,
                })
                analyzed += 1
            except Exception as e:
                print(f"[Batch Analysis Error] {obj['id']}: {e}")
                results.append({"id": obj["id"], "title": obj["title"], "success": False, "error": str(e)})
                errors += 1

        yield f"data: {json.dumps({'type': 'done', 'success': True, 'analyzed': analyzed, 'errors': errors, 'results': results})}\n\n"

    return generate


# ============= ANALYTICS / DÉTAILS / HISTORIQUE =============

def get_analytics(current_user: dict, entity=None, manager=None, period=None) -> dict:
    """Tableaux de bord analytiques RH."""
    _require_hr(current_user)
    with db_session() as db:
        return analysis_repository.hr_analytics(db, entity, manager, period)


def get_objectives_details(
    current_user: dict, entity=None, manager=None, period=None,
    smart_filter=None, alignment_filter=None,
) -> dict:
    """Liste détaillée des objectifs RH avec filtres."""
    _require_hr(current_user)
    with db_session() as db:
        rows = analysis_repository.hr_objectives_details(
            db, entity, manager, period, smart_filter, alignment_filter
        )
    for obj in rows:
        for key in ("smart_analysis", "alignment_analysis"):
            if obj.get(key):
                try:
                    obj[key] = json.loads(obj[key])
                except (json.JSONDecodeError, TypeError):
                    obj[key] = None
    return {"objectives": rows, "total": len(rows)}


def get_imports(current_user: dict) -> dict:
    """Historique des imports RH."""
    _require_hr(current_user)
    with db_session() as db:
        return {"imports": [i.to_dict() for i in hr_import_repository.list_recent(db, limit=50)]}


# ============= EXPORTS (délégués à export_service) =============

_CRITERIA = [
    ("S - Spécifique", "s_score"),
    ("M - Mesurable", "m_score"),
    ("A - Atteignable", "a_score"),
    ("R - Réaliste", "r_score"),
    ("T - Temporel", "t_score"),
]


def export_excel(current_user: dict, entity=None, manager=None, period=None):
    """Construit le rapport Excel multi-feuilles RH."""
    _require_hr(current_user)
    with db_session() as db:
        detailed = analysis_repository.hr_export_detailed(db, entity, manager, period)
        entity_stats = analysis_repository.hr_export_entity_stats(db, entity, manager, period)
        manager_stats = analysis_repository.hr_export_manager_stats(db, entity, manager, period)
        total_obj = analysis_repository.hr_count_objectives(db, entity, manager, period) or 1
        failures = [
            (label, analysis_repository.hr_criteria_failures(db, col, entity, manager, period))
            for label, col in _CRITERIA
        ]
    return export_service.generate_hr_excel(detailed, entity_stats, manager_stats, total_obj, failures)


def export_pdf(current_user: dict, entity=None, manager=None, period=None):
    """Construit le rapport PDF synthétique RH."""
    _require_hr(current_user)
    with db_session() as db:
        objectives = analysis_repository.hr_export_pdf_rows(db, entity, manager, period)
        total = analysis_repository.hr_count_objectives(db, entity, manager, period)
        smart_count = analysis_repository.hr_pdf_smart_count(db, entity, manager, period)
        avg_alignment = analysis_repository.hr_pdf_avg_alignment(db, entity, manager, period)
    return export_service.generate_hr_pdf(objectives, total, smart_count, avg_alignment)
