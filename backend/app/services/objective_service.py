#!/usr/bin/env python3
"""
services/objective_service.py — Logique métier des objectifs.

RÔLE
    Gère les objectifs personnels (CRUD filtré par rôle), le catalogue stratégique
    d'entreprise et les statistiques d'équipe génériques. Toute la persistance
    passe par les repositories ; le contrôle d'accès par le permission_service.

EMPLACEMENT MVC
    Couche "services" : logique métier.
"""

import io
import json
import logging
import secrets
from datetime import datetime, timezone
from typing import List

import openpyxl
from fastapi.responses import RedirectResponse, StreamingResponse
from openpyxl.styles import Font, PatternFill

logger = logging.getLogger(__name__)

from fastapi import HTTPException

from app.database.session import db_session
from app.repositories import objective_repository, company_objective_repository
from app.services import permission_service as perm

_EMPTY_DISTRIBUTION = {"not_started": 0, "in_progress": 0, "completed": 0, "on_hold": 0}


def _parse_analyses(obj: dict) -> dict:
    """Désérialise les analyses JSON (smart/alignment) stockées en TEXT."""
    for key in ("smart_analysis", "alignment_analysis"):
        if obj.get(key):
            try:
                obj[key] = json.loads(obj[key])
            except (json.JSONDecodeError, TypeError):
                obj[key] = None
    return obj


# ============= OBJECTIFS PERSONNELS =============

def list_objectives(current_user: dict) -> dict:
    """Liste les objectifs accessibles selon le rôle."""
    accessible_ids = perm.get_accessible_user_ids(current_user)
    if not accessible_ids:
        return {"objectives": []}

    with db_session() as db:
        rows = objective_repository.list_by_user_ids(db, accessible_ids)
        objectives = []
        for o in rows:
            obj = o.to_dict()
            obj["user_name"] = o.user.name if o.user else None
            obj["user_email"] = o.user.email if o.user else None
            objectives.append(_parse_analyses(obj))
        return {"objectives": objectives}


def create_objective(current_user: dict, request) -> dict:
    """Crée un objectif (uniquement pour soi-même)."""
    obj_id = f"obj-{secrets.token_hex(6)}"
    with db_session() as db:
        objective_repository.create(
            db,
            id=obj_id,
            user_id=current_user["id"],
            title=request.title,
            description=request.description,
            context=request.context,
        )
    return {
        "success": True,
        "objective": {
            "id": obj_id,
            "user_id": current_user["id"],
            "title": request.title,
            "description": request.description,
            "status": "draft",
            "progress": "not_started",
        },
    }


def get_objective(objective_id: str, current_user: dict) -> dict:
    """Récupère un objectif avec contrôle d'accès RBAC."""
    with db_session() as db:
        o = objective_repository.get_with_user(db, objective_id)
        if not o:
            raise HTTPException(status_code=404, detail="Objectif non trouvé.")
        obj = o.to_dict()
        obj["user_name"] = o.user.name if o.user else None

        if current_user["role"] == "COLLAB":
            perm.check_collab_access(current_user, obj["user_id"])
        elif current_user["role"] == "MANAGER":
            perm.check_manager_access(current_user, obj["user_id"])

        return {"objective": _parse_analyses(obj)}


def update_objective(objective_id: str, request, current_user: dict) -> dict:
    """Met à jour un objectif (propriétaire ou HR uniquement)."""
    with db_session() as db:
        o = objective_repository.get(db, objective_id)
        if not o:
            raise HTTPException(status_code=404, detail="Objectif non trouvé.")
        if current_user["role"] != "HR" and o.user_id != current_user["id"]:
            raise HTTPException(status_code=403, detail="Vous ne pouvez modifier que vos propres objectifs.")

        values: dict = {}
        if request.title is not None:
            values["title"] = request.title
        if request.description is not None:
            values["description"] = request.description
        if request.status is not None:
            values["status"] = request.status
        if request.progress is not None:
            values["progress"] = request.progress
        if request.smart_analysis is not None:
            values["smart_analysis"] = json.dumps(request.smart_analysis)
        if request.alignment_analysis is not None:
            values["alignment_analysis"] = json.dumps(request.alignment_analysis)

        if values:
            values["updated_at"] = datetime.now(timezone.utc).isoformat()
            objective_repository.update_fields(db, objective_id, values)
        else:
            logger.debug("update_objective: aucun champ fourni pour %s, rien à faire.", objective_id)

    return {"success": True}


def delete_objective(objective_id: str, current_user: dict) -> dict:
    """Supprime un objectif (propriétaire ou HR uniquement)."""
    with db_session() as db:
        o = objective_repository.get(db, objective_id)
        if not o:
            raise HTTPException(status_code=404, detail="Objectif non trouvé.")
        if current_user["role"] != "HR" and o.user_id != current_user["id"]:
            raise HTTPException(status_code=403, detail="Vous ne pouvez supprimer que vos propres objectifs.")
        objective_repository.delete_by_id(db, objective_id)
    return {"success": True}


# ============= OBJECTIFS STRATÉGIQUES (CATALOGUE ENTREPRISE) =============

def list_company_objectives() -> dict:
    """Liste le catalogue stratégique (lecture publique)."""
    with db_session() as db:
        return {"company_objectives": [c.to_dict() for c in company_objective_repository.list_all(db)]}


def create_company_objective(request, current_user: dict) -> dict:
    """Crée un objectif stratégique (HR ou Manager seulement)."""
    _require_manager_or_hr(current_user)

    obj_id = request.id or f"STRAT-{secrets.token_hex(4)}"
    with db_session() as db:
        company_objective_repository.create(
            db,
            id=obj_id,
            title=request.title,
            description=request.description,
            department=request.department,
            created_by=current_user["id"],
        )
    return {"success": True, "id": obj_id}


def batch_create_company_objectives(objectives: List, current_user: dict) -> dict:
    """Import batch d'objectifs stratégiques (HR ou Manager seulement)."""
    _require_manager_or_hr(current_user)

    created = 0
    errors = []
    with db_session() as db:
        for obj in objectives:
            obj_id = obj.id or f"STRAT-{secrets.token_hex(4)}"
            try:
                company_objective_repository.merge(db, {
                    "id": obj_id,
                    "title": obj.title,
                    "description": obj.description,
                    "department": obj.department,
                    "created_by": current_user["id"],
                })
                created += 1
            except Exception as exc:
                logger.warning("Objectif ignoré lors de l'import batch (id=%s) : %s", obj_id, exc)
                errors.append({"id": obj_id, "title": obj.title, "error": str(exc)})
    return {"success": True, "created": created, "errors": errors}


def _require_manager_or_hr(current_user: dict):
    if current_user["role"] not in ("MANAGER", "HR"):
        raise HTTPException(status_code=403,
                            detail="Seuls les managers et RH peuvent gérer les objectifs stratégiques.")


def update_company_objective(objective_id: str, request, current_user: dict) -> dict:
    """Met à jour un objectif stratégique (HR ou Manager)."""
    _require_manager_or_hr(current_user)
    values = {}
    if request.title is not None:
        values["title"] = request.title.strip()
    if request.description is not None:
        values["description"] = request.description.strip()
    if request.department is not None:
        values["department"] = request.department.strip()
    with db_session() as db:
        updated = company_objective_repository.update(db, objective_id, values)
        if updated is None:
            raise HTTPException(status_code=404, detail="Objectif stratégique introuvable.")
    return {"success": True, "id": objective_id}


def delete_company_objective(objective_id: str, current_user: dict) -> dict:
    """Supprime un objectif stratégique (HR ou Manager)."""
    _require_manager_or_hr(current_user)
    with db_session() as db:
        deleted = company_objective_repository.delete_by_id(db, objective_id)
        if not deleted:
            raise HTTPException(status_code=404, detail="Objectif stratégique introuvable.")
    return {"success": True}


def export_company_objectives_excel(current_user: dict):
    """Exporte le catalogue stratégique dans un modèle Excel (ré-importable)."""
    _require_manager_or_hr(current_user)

    with db_session() as db:
        objs = company_objective_repository.list_all(db)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Objectifs strategiques"
    headers = ["ID", "Titre", "Description", "Departement", "Priorite"]
    ws.append(headers)
    head_font = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="7C3AED")
    for cell in ws[1]:
        cell.font = head_font
        cell.fill = head_fill
    if objs:
        for o in objs:
            ws.append([o.id, o.title, o.description or "", o.department or "", o.priority or 1])
    else:
        # Modèle vide : deux lignes d'exemple pour guider la saisie.
        ws.append(["AXE-01", "Transformation digitale", "Description de l'axe...", "Direction", 1])
        ws.append(["", "Nouvel axe (ID laisse vide = genere automatiquement)", "", "", 1])
    for col, width in zip("ABCDE", (12, 40, 55, 20, 10)):
        ws.column_dimensions[col].width = width

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    filename = f"objectifs_strategiques_{datetime.now().strftime('%Y%m%d')}.xlsx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


# ============= STATISTIQUES D'ÉQUIPE (génériques) =============

def compute_team_stats(team_ids: List[str]) -> dict:
    """
    Cœur de calcul des statistiques d'équipe (réutilisé par les domaines
    génériques et manager). Retourne total, taux SMART, alignement moyen et
    répartition d'avancement.
    """
    if not team_ids:
        return {
            "total_objectives": 0,
            "smart_rate": None,
            "avg_alignment": None,
            "progress_distribution": dict(_EMPTY_DISTRIBUTION),
        }

    with db_session() as db:
        total = objective_repository.count_by_user_ids(db, team_ids)
        if total == 0:
            return {
                "total_objectives": 0,
                "smart_rate": None,
                "avg_alignment": None,
                "progress_distribution": dict(_EMPTY_DISTRIBUTION),
            }

        smart_count = 0
        total_analyzed = 0
        for raw in objective_repository.get_smart_analyses(db, team_ids):
            try:
                analysis = json.loads(raw)
            except (json.JSONDecodeError, TypeError):
                continue
            total_analyzed += 1
            if analysis.get("overall"):
                smart_count += 1
        smart_rate = round((smart_count / total_analyzed) * 100) if total_analyzed > 0 else None

        alignment_scores = []
        for raw in objective_repository.get_alignment_analyses(db, team_ids):
            try:
                analysis = json.loads(raw)
            except (json.JSONDecodeError, TypeError):
                continue
            if "alignment_score" in analysis:
                alignment_scores.append(analysis["alignment_score"])
        avg_alignment = round(sum(alignment_scores) / len(alignment_scores)) if alignment_scores else None

        progress_dist = dict(_EMPTY_DISTRIBUTION)
        for progress, count in objective_repository.get_progress_counts(db, team_ids):
            if progress in progress_dist:
                progress_dist[progress] = count

        return {
            "total_objectives": total,
            "smart_rate": smart_rate,
            "avg_alignment": avg_alignment,
            "progress_distribution": progress_dist,
        }


def get_team_stats_generic(current_user: dict) -> dict:
    """
    Statistiques d'équipe génériques (endpoint /api/team/stats) :
      - COLLAB  : accès refusé
      - HR      : redirigé vers /api/hr/analytics
      - MANAGER : stats de son équipe
    """
    if current_user["role"] == "COLLAB":
        raise HTTPException(status_code=403, detail="Accès réservé aux managers.")
    if current_user["role"] == "HR":
        return RedirectResponse(url="/api/hr/analytics", status_code=303)

    team_ids = perm.get_manager_team_ids(current_user["id"])
    stats = compute_team_stats(team_ids)
    stats["team_members"] = len(team_ids)
    stats["scope"] = "manager_team"
    return stats
