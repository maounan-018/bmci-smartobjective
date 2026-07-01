#!/usr/bin/env python3
"""
services/export_service.py — Génération de fichiers d'export (Excel / PDF).

RÔLE
    Centralise la détection des bibliothèques optionnelles (openpyxl, reportlab)
    et la génération des rapports Manager (Excel et PDF). Les controllers se
    contentent d'appeler ces fonctions avec les données déjà filtrées par rôle.

EMPLACEMENT MVC
    Couche "services" : génération de documents, sans route HTTP.
"""

import io
import logging
from datetime import datetime

from fastapi import HTTPException
from fastapi.responses import StreamingResponse

from app.utils.helpers import safe_int, safe_float, is_smart_valid

logger = logging.getLogger(__name__)

# --- Détection du support Excel (openpyxl) ---
try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    EXCEL_SUPPORT = True
except Exception as e:  # pragma: no cover
    EXCEL_SUPPORT = False
    logger.warning("openpyxl indisponible - export Excel limité (%s)", e)

# --- Détection du support PDF (reportlab) ---
try:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer, PageBreak
    from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY
    PDF_SUPPORT = True
except Exception as e:  # pragma: no cover
    PDF_SUPPORT = False
    logger.warning("reportlab indisponible - export PDF limité (%s)", e)


def generate_manager_excel(objectives, personal_analyses, current_user) -> StreamingResponse:
    """Génère le rapport Excel du Manager (objectifs équipe + analyses perso + résumé)."""
    if not EXCEL_SUPPORT:
        raise HTTPException(status_code=500, detail="openpyxl non installé. Exécutez: pip install openpyxl")

    wb = Workbook()

    # Styles
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="059669", end_color="059669", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    thin_border = Border(
        left=Side(style='thin'), right=Side(style='thin'),
        top=Side(style='thin'), bottom=Side(style='thin')
    )

    # === FEUILLE 1: Objectifs de l'équipe ===
    ws1 = wb.active
    ws1.title = "Objectifs Équipe"

    headers1 = ["Collaborateur", "Email", "Titre Objectif", "Description", "Statut SMART",
                "Score SMART %", "S", "M", "A", "R", "T", "Alignement %", "Avancement", "Date Création"]

    for col, header in enumerate(headers1, 1):
        cell = ws1.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = thin_border

    for row_idx, obj in enumerate(objectives, 2):
        s_score = obj.get('s_score') or 0
        m_score = obj.get('m_score') or 0
        a_score = obj.get('a_score') or 0
        r_score = obj.get('r_score') or 0
        t_score = obj.get('t_score') or 0
        is_smart = (s_score + m_score + a_score + r_score + t_score) >= 4

        user_name = obj.get('user_name') or ''
        email = obj.get('email') or ''
        title = obj.get('title') or ''
        description = obj.get('description') or ''
        progress = obj.get('progress') or ''
        created_at = obj.get('created_at') or ''

        values = [
            user_name, email, title, description,
            'SMART' if is_smart else 'NON SMART',
            obj.get('smart_score') or 0,
            '✓' if s_score else '✗',
            '✓' if m_score else '✗',
            '✓' if a_score else '✗',
            '✓' if r_score else '✗',
            '✓' if t_score else '✗',
            obj.get('alignment_score') or 0,
            progress.replace('_', ' ').upper() if progress else '',
            created_at[:10] if created_at else ''
        ]
        for col, val in enumerate(values, 1):
            cell = ws1.cell(row=row_idx, column=col, value=val)
            cell.border = thin_border
            if col in [6, 12]:  # Colonnes de score
                cell.alignment = Alignment(horizontal="center")

    ws1.column_dimensions['A'].width = 20
    ws1.column_dimensions['C'].width = 35
    ws1.column_dimensions['D'].width = 40

    # === FEUILLE 2: Mes Analyses Personnelles ===
    ws2 = wb.create_sheet("Mes Analyses SMART")

    headers2 = ["Date", "Titre Objectif", "Score SMART %", "S", "M", "A", "R", "T",
                "Alignement %", "Nb Issues", "Issue Principale"]

    for col, header in enumerate(headers2, 1):
        cell = ws2.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = thin_border

    for row_idx, analysis in enumerate(personal_analyses, 2):
        obj_title = analysis.get('objective_title') or ''
        top_issue = analysis.get('top_issue') or ''
        values = [
            analysis.get('analyzed_at', '')[:16].replace('T', ' ') if analysis.get('analyzed_at') else '',
            obj_title,
            round(safe_float(analysis.get('smart_score'))),
            '✓' if safe_int(analysis.get('s_score')) else '✗',
            '✓' if safe_int(analysis.get('m_score')) else '✗',
            '✓' if safe_int(analysis.get('a_score')) else '✗',
            '✓' if safe_int(analysis.get('r_score')) else '✗',
            '✓' if safe_int(analysis.get('t_score')) else '✗',
            safe_float(analysis.get('alignment_score')),
            safe_int(analysis.get('issues_count')),
            top_issue[:50] + '...' if len(top_issue) > 50 else top_issue
        ]
        for col, val in enumerate(values, 1):
            cell = ws2.cell(row=row_idx, column=col, value=val)
            cell.border = thin_border

    ws2.column_dimensions['B'].width = 40
    ws2.column_dimensions['K'].width = 50

    # === FEUILLE 3: Résumé ===
    ws3 = wb.create_sheet("Résumé")

    total_obj = len(objectives)
    smart_count = sum(1 for o in objectives if is_smart_valid(o))

    summary_data = [
        ["RAPPORT D'ANALYSE SMART - ÉQUIPE", ""],
        ["", ""],
        ["Manager", current_user.get('name', '')],
        ["Date du rapport", datetime.now().strftime('%d/%m/%Y %H:%M')],
        ["", ""],
        ["STATISTIQUES ÉQUIPE", ""],
        ["Total objectifs", total_obj],
        ["Objectifs SMART", smart_count],
        ["Objectifs Non SMART", total_obj - smart_count],
        ["Taux de conformité SMART", f"{round(smart_count/total_obj*100) if total_obj else 0}%"],
        ["", ""],
        ["MES ANALYSES PERSONNELLES", ""],
        ["Total analyses effectuées", len(personal_analyses)],
    ]

    for row_idx, (label, value) in enumerate(summary_data, 1):
        cell_label = ws3.cell(row=row_idx, column=1, value=label)
        cell_value = ws3.cell(row=row_idx, column=2, value=value)
        if label in ["RAPPORT D'ANALYSE SMART - ÉQUIPE", "STATISTIQUES ÉQUIPE", "MES ANALYSES PERSONNELLES"]:
            cell_label.font = Font(bold=True, size=14)
        cell_label.border = thin_border
        cell_value.border = thin_border

    ws3.column_dimensions['A'].width = 35
    ws3.column_dimensions['B'].width = 25

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    filename = f"Rapport_Manager_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"

    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )


def generate_manager_pdf(objectives, personal_analyses, current_user) -> StreamingResponse:
    """Génère le rapport PDF du Manager (résumé + objectifs équipe + analyses perso)."""
    if not PDF_SUPPORT:
        raise HTTPException(status_code=500, detail="reportlab non installé. Exécutez: pip install reportlab")

    output = io.BytesIO()
    doc = SimpleDocTemplate(output, pagesize=landscape(A4),
                            leftMargin=1*cm, rightMargin=1*cm,
                            topMargin=1.5*cm, bottomMargin=1.5*cm)

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle('Title', parent=styles['Heading1'], fontSize=18,
                                 textColor=colors.HexColor('#059669'), alignment=TA_CENTER)
    subtitle_style = ParagraphStyle('Subtitle', parent=styles['Heading2'], fontSize=14,
                                    textColor=colors.HexColor('#334155'))
    normal_style = styles['Normal']

    elements = []

    elements.append(Paragraph("RAPPORT D'ANALYSE SMART", title_style))
    elements.append(Spacer(1, 0.5*cm))
    elements.append(Paragraph(f"Manager: {current_user.get('name', '')}", normal_style))
    elements.append(Paragraph(f"Date: {datetime.now().strftime('%d/%m/%Y %H:%M')}", normal_style))
    elements.append(Spacer(1, 1*cm))

    total_obj = len(objectives)
    smart_count = sum(1 for o in objectives if is_smart_valid(o))

    elements.append(Paragraph("RÉSUMÉ", subtitle_style))
    summary_data = [
        ["Total Objectifs Équipe", str(total_obj)],
        ["Objectifs SMART", str(smart_count)],
        ["Objectifs Non SMART", str(total_obj - smart_count)],
        ["Taux de conformité", f"{round(smart_count/total_obj*100) if total_obj else 0}%"],
        ["Analyses personnelles", str(len(personal_analyses))]
    ]
    summary_table = Table(summary_data, colWidths=[8*cm, 4*cm])
    summary_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor('#f1f5f9')),
        ('TEXTCOLOR', (0, 0), (-1, -1), colors.HexColor('#334155')),
        ('FONTNAME', (0, 0), (-1, -1), 'Helvetica'),
        ('FONTSIZE', (0, 0), (-1, -1), 10),
        ('ALIGN', (1, 0), (1, -1), 'CENTER'),
        ('FONTNAME', (1, 0), (1, -1), 'Helvetica-Bold'),
        ('BOX', (0, 0), (-1, -1), 1, colors.HexColor('#059669')),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#cbd5e1')),
    ]))
    elements.append(summary_table)
    elements.append(Spacer(1, 1*cm))

    if objectives:
        elements.append(Paragraph("OBJECTIFS DE L'ÉQUIPE", subtitle_style))
        elements.append(Spacer(1, 0.3*cm))

        table_data = [["Collaborateur", "Titre", "Score", "S", "M", "A", "R", "T", "Statut"]]

        for obj in objectives[:30]:  # Limiter à 30 lignes pour le PDF
            user_name = obj.get('user_name') or ''
            obj_title = obj.get('title') or ''
            table_data.append([
                user_name[:20],
                obj_title[:40] + ('...' if len(obj_title) > 40 else ''),
                f"{round(safe_float(obj.get('smart_score')))}%",
                '✓' if safe_int(obj.get('s_score')) else '✗',
                '✓' if safe_int(obj.get('m_score')) else '✗',
                '✓' if safe_int(obj.get('a_score')) else '✗',
                '✓' if safe_int(obj.get('r_score')) else '✗',
                '✓' if safe_int(obj.get('t_score')) else '✗',
                'SMART' if is_smart_valid(obj) else 'NON SMART'
            ])

        obj_table = Table(table_data, colWidths=[3.5*cm, 7*cm, 1.5*cm, 1*cm, 1*cm, 1*cm, 1*cm, 1*cm, 2.5*cm])
        obj_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#059669')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, -1), 8),
            ('ALIGN', (2, 0), (-1, -1), 'CENTER'),
            ('BOX', (0, 0), (-1, -1), 1, colors.HexColor('#059669')),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#e2e8f0')),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f8fafc')]),
        ]))
        elements.append(obj_table)

    if personal_analyses:
        elements.append(PageBreak())
        elements.append(Paragraph("MES ANALYSES PERSONNELLES", subtitle_style))
        elements.append(Spacer(1, 0.3*cm))

        pa_data = [["Date", "Titre", "Score", "S", "M", "A", "R", "T"]]
        for analysis in personal_analyses[:20]:
            title = analysis.get('objective_title') or ''
            pa_data.append([
                analysis.get('analyzed_at', '')[:10] if analysis.get('analyzed_at') else '',
                title[:50] + ('...' if len(title) > 50 else ''),
                f"{round(safe_float(analysis.get('smart_score')))}%",
                '✓' if safe_int(analysis.get('s_score')) else '✗',
                '✓' if safe_int(analysis.get('m_score')) else '✗',
                '✓' if safe_int(analysis.get('a_score')) else '✗',
                '✓' if safe_int(analysis.get('r_score')) else '✗',
                '✓' if safe_int(analysis.get('t_score')) else '✗',
            ])

        pa_table = Table(pa_data, colWidths=[2.5*cm, 10*cm, 1.5*cm, 1*cm, 1*cm, 1*cm, 1*cm, 1*cm])
        pa_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#7c3aed')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, -1), 8),
            ('ALIGN', (2, 0), (-1, -1), 'CENTER'),
            ('BOX', (0, 0), (-1, -1), 1, colors.HexColor('#7c3aed')),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#e2e8f0')),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#faf5ff')]),
        ]))
        elements.append(pa_table)

    doc.build(elements)
    output.seek(0)

    filename = f"Rapport_Manager_{datetime.now().strftime('%Y%m%d_%H%M%S')}.pdf"

    return StreamingResponse(
        output,
        media_type="application/pdf",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )


# ============================================================================
# RAPPORTS RH (données pré-extraites par hr_service via les repositories)
# ============================================================================

def generate_hr_excel(detailed_rows, entity_stats, manager_stats, total_obj, failures) -> StreamingResponse:
    """
    Rapport Excel RH multi-feuilles : objectifs détaillés, stats par entité,
    stats par manager, points bloquants SMART.

    Paramètres (déjà filtrés/agrégés au niveau repository) :
        detailed_rows : lignes (tuples) dans l'ordre des colonnes d'export.
        entity_stats / manager_stats : tuples (libellé, count, avg_smart, avg_alignment).
        total_obj     : total d'objectifs du périmètre (>= 1).
        failures      : liste (libellé_critère, nb_défaillances).
    """
    if not EXCEL_SUPPORT:
        raise HTTPException(status_code=500, detail="openpyxl non installé.")

    wb = Workbook()
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="059669", end_color="059669", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    thin_border = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"), bottom=Side(style="thin"),
    )

    def _write_headers(ws, headers):
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
            cell.border = thin_border

    # --- Feuille 1 : objectifs détaillés ---
    ws1 = wb.active
    ws1.title = "Objectifs Détaillés"
    _write_headers(ws1, [
        "ID", "Collaborateur", "Email", "Entité", "Manager", "Période",
        "Titre", "Description", "Score SMART (%)", "S", "M", "A", "R", "T",
        "Score Alignement (%)", "Nb Issues", "Issue Principale", "Date Création",
    ])
    for row_idx, row in enumerate(detailed_rows, 2):
        for col_idx, value in enumerate(row, 1):
            cell = ws1.cell(row=row_idx, column=col_idx, value=value)
            cell.border = thin_border
            if col_idx in [9, 15]:
                cell.number_format = "0.0"
                if value is not None:
                    if value >= 80:
                        cell.fill = PatternFill(start_color="D1FAE5", end_color="D1FAE5", fill_type="solid")
                    elif value >= 50:
                        cell.fill = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
                    else:
                        cell.fill = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
    for col in ws1.columns:
        max_length = 0
        column = col[0].column_letter
        for cell in col:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except Exception:
                pass
        ws1.column_dimensions[column].width = min(max_length + 2, 50)

    # --- Feuille 2 : stats par entité ---
    ws2 = wb.create_sheet(title="Stats par Entité")
    _write_headers(ws2, ["Entité", "Nb Objectifs", "Score SMART Moyen", "Score Alignement Moyen"])
    for row_idx, row in enumerate(entity_stats, 2):
        for col_idx, value in enumerate(row, 1):
            cell = ws2.cell(row=row_idx, column=col_idx, value=value)
            cell.border = thin_border
            if col_idx in [3, 4]:
                cell.number_format = "0.0"

    # --- Feuille 3 : stats par manager ---
    ws3 = wb.create_sheet(title="Stats par Manager")
    _write_headers(ws3, ["Manager", "Nb Objectifs", "Score SMART Moyen", "Score Alignement Moyen"])
    for row_idx, row in enumerate(manager_stats, 2):
        for col_idx, value in enumerate(row, 1):
            cell = ws3.cell(row=row_idx, column=col_idx, value=value)
            cell.border = thin_border
            if col_idx in [3, 4]:
                cell.number_format = "0.0"

    # --- Feuille 4 : points bloquants ---
    ws4 = wb.create_sheet(title="Points Bloquants")
    _write_headers(ws4, ["Critère SMART", "Nb Défaillances", "% Défaillance"])
    for row_idx, (criteria, failures_count) in enumerate(failures, 2):
        ws4.cell(row=row_idx, column=1, value=criteria).border = thin_border
        ws4.cell(row=row_idx, column=2, value=int(failures_count or 0)).border = thin_border
        pct_cell = ws4.cell(row=row_idx, column=3, value=(failures_count or 0) / total_obj if total_obj else 0)
        pct_cell.border = thin_border
        pct_cell.number_format = "0.0%"

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    filename = f"Rapport_RH_BMCI_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )


def generate_hr_pdf(objectives, total, smart_count, avg_alignment) -> StreamingResponse:
    """Rapport PDF RH : synthèse globale + détail des objectifs (limité)."""
    if not PDF_SUPPORT:
        raise HTTPException(status_code=500, detail="reportlab non installé. Exécutez: pip install reportlab")

    output = io.BytesIO()
    doc = SimpleDocTemplate(output, pagesize=landscape(A4),
                            leftMargin=1 * cm, rightMargin=1 * cm,
                            topMargin=1.5 * cm, bottomMargin=1.5 * cm)
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle("Title", parent=styles["Heading1"], fontSize=20,
                                 textColor=colors.HexColor("#059669"), alignment=TA_CENTER)
    subtitle_style = ParagraphStyle("Subtitle", parent=styles["Heading2"], fontSize=14,
                                    textColor=colors.HexColor("#334155"))
    elements = []

    elements.append(Paragraph("RAPPORT D'ANALYSE RH - OBJECTIFS SMART", title_style))
    elements.append(Spacer(1, 0.5 * cm))
    elements.append(Paragraph(f"BMCI - {datetime.now().strftime('%d/%m/%Y %H:%M')}", styles["Normal"]))
    elements.append(Spacer(1, 1 * cm))

    elements.append(Paragraph("SYNTHÈSE GLOBALE", subtitle_style))
    summary_data = [
        ["Total Objectifs", str(total)],
        ["Objectifs SMART (≥80%)", str(smart_count)],
        ["Taux de conformité", f"{round(smart_count / total * 100) if total else 0}%"],
        ["Alignement moyen", f"{round(avg_alignment)}%"],
    ]
    summary_table = Table(summary_data, colWidths=[8 * cm, 4 * cm])
    summary_table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#f1f5f9")),
        ("TEXTCOLOR", (0, 0), (-1, -1), colors.HexColor("#334155")),
        ("FONTNAME", (0, 0), (-1, -1), "Helvetica"),
        ("FONTSIZE", (0, 0), (-1, -1), 11),
        ("ALIGN", (1, 0), (1, -1), "CENTER"),
        ("FONTNAME", (1, 0), (1, -1), "Helvetica-Bold"),
        ("BOX", (0, 0), (-1, -1), 1, colors.HexColor("#059669")),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#cbd5e1")),
    ]))
    elements.append(summary_table)
    elements.append(Spacer(1, 1 * cm))

    elements.append(Paragraph("DÉTAIL DES OBJECTIFS", subtitle_style))
    elements.append(Spacer(1, 0.3 * cm))
    table_data = [["Collaborateur", "Entité", "Titre", "Score", "S", "M", "A", "R", "T", "Align."]]
    for obj in objectives:
        user_name = obj.get("user_name") or ""
        entity = obj.get("entity") or ""
        title = obj.get("title") or ""
        table_data.append([
            user_name[:15],
            entity[:10],
            title[:35] + ("..." if len(title) > 35 else ""),
            f"{round(obj.get('smart_score') or 0)}%",
            "✓" if obj.get("s_score") else "✗",
            "✓" if obj.get("m_score") else "✗",
            "✓" if obj.get("a_score") else "✗",
            "✓" if obj.get("r_score") else "✗",
            "✓" if obj.get("t_score") else "✗",
            f"{round(obj.get('alignment_score') or 0)}%",
        ])
    obj_table = Table(table_data, colWidths=[3 * cm, 2.5 * cm, 7 * cm, 1.5 * cm, 1 * cm, 1 * cm, 1 * cm, 1 * cm, 1 * cm, 1.5 * cm])
    obj_table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#059669")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("ALIGN", (3, 0), (-1, -1), "CENTER"),
        ("BOX", (0, 0), (-1, -1), 1, colors.HexColor("#059669")),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#e2e8f0")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f8fafc")]),
    ]))
    elements.append(obj_table)

    doc.build(elements)
    output.seek(0)
    filename = f"Rapport_RH_BMCI_{datetime.now().strftime('%Y%m%d_%H%M%S')}.pdf"
    return StreamingResponse(
        output,
        media_type="application/pdf",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )


def _esc(s) -> str:
    """Échappe le texte destiné à un Paragraph reportlab (XML)."""
    return str(s or "").replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def generate_hr_comex_pdf(data, paragraphs, synthesis, period_label="Campagne en cours") -> StreamingResponse:
    """Rapport COMEX (Direction) : photo de la campagne, KPI + paragraphes
    observationnels. AUCUNE comparaison (personnes, entités, périodes), AUCUN nom.

    `data` = agrégats de hr_analysis_service.get_analytics.
    `paragraphs` = dict clé->texte IA observationnel (repli factuel si vide).
    `synthesis` = synthèse générale IA (repli factuel si vide).
    """
    if not PDF_SUPPORT:
        raise HTTPException(status_code=500, detail="reportlab non installé. Exécutez: pip install reportlab")

    EMERALD = colors.HexColor("#059669")
    SLATE = colors.HexColor("#334155")
    GREY = colors.HexColor("#64748b")
    LIGHT = colors.HexColor("#f1f5f9")
    ROSE = colors.HexColor("#e11d48")
    VIOLET = colors.HexColor("#7c3aed")

    styles = getSampleStyleSheet()
    h1 = ParagraphStyle("ComexH1", parent=styles["Heading1"], fontSize=16, textColor=EMERALD, spaceAfter=1)
    h1sub = ParagraphStyle("ComexH1sub", parent=styles["Normal"], fontSize=9, textColor=GREY, spaceAfter=10)
    kpilbl = ParagraphStyle("ComexLbl", parent=styles["Normal"], fontSize=9, textColor=GREY, spaceAfter=3)
    body = ParagraphStyle("ComexBody", parent=styles["Normal"], fontSize=9.5, textColor=SLATE,
                          leading=13, alignment=TA_JUSTIFY, spaceAfter=4)
    small = ParagraphStyle("ComexSmall", parent=styles["Normal"], fontSize=8, textColor=GREY)
    cover_title = ParagraphStyle("ComexCover", parent=styles["Heading1"], fontSize=22, textColor=EMERALD,
                                 alignment=TA_CENTER, leading=26)
    cover_sub = ParagraphStyle("ComexCoverSub", parent=styles["Normal"], fontSize=12, textColor=SLATE,
                               alignment=TA_CENTER)

    avg = data.get("avg_scores", {}) or {}
    sr = data.get("smart_rate", {}) or {}
    vol = data.get("volumetry", {}) or {}
    lc = data.get("low_confidence", {}) or {}
    cov = data.get("axis_coverage", {}) or {}
    thr = data.get("thresholds", {}) or {}
    axes = cov.get("axes") or []

    def pctv(x):
        return f"{round(x or 0)}%"

    to_adjust = max(0, (sr.get("total", 0) or 0) - (sr.get("ok", 0) or 0))
    orphans = len([a for a in axes if a.get("orphan")])
    over = len([a for a in axes if a.get("over_concentrated")])

    def para(key, fallback):
        txt = (paragraphs or {}).get(key) or ""
        return Paragraph(_esc(txt or fallback), body)

    def kpi_cell(num, label, key, fallback, color=SLATE):
        numstyle = ParagraphStyle(f"num_{key}", parent=styles["Normal"], fontSize=21,
                                  textColor=color, leading=23)
        return [Paragraph(_esc(num), numstyle), Paragraph(_esc(label), kpilbl),
                para(key, fallback)]

    def table(rows, widths, header=EMERALD):
        t = Table(rows, colWidths=widths)
        t.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), header),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, -1), 9),
            ("ALIGN", (1, 0), (-1, -1), "CENTER"),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#e2e8f0")),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, LIGHT]),
            ("TOPPADDING", (0, 0), (-1, -1), 5),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
            ("LEFTPADDING", (0, 0), (-1, -1), 6),
        ]))
        return t

    output = io.BytesIO()
    doc = SimpleDocTemplate(output, pagesize=A4, leftMargin=2 * cm, rightMargin=2 * cm,
                            topMargin=2 * cm, bottomMargin=2 * cm)
    elements = []

    # ===== PAGE 1 — Couverture =====
    elements.append(Spacer(1, 4.5 * cm))
    elements.append(Paragraph("Rapport de pilotage des objectifs", cover_title))
    elements.append(Spacer(1, 0.4 * cm))
    elements.append(Paragraph("Synthèse SMART &amp; alignement stratégique — Direction / COMEX", cover_sub))
    elements.append(Spacer(1, 1.2 * cm))
    cover = Table([
        ["Période", period_label],
        ["Périmètre", "Groupe BMCI"],
        ["Édité le", datetime.now().strftime("%d/%m/%Y")],
        ["Préparé par", "Direction des Ressources Humaines"],
    ], colWidths=[5 * cm, 9 * cm])
    cover.setStyle(TableStyle([
        ("FONTSIZE", (0, 0), (-1, -1), 11),
        ("TEXTCOLOR", (0, 0), (0, -1), GREY),
        ("TEXTCOLOR", (1, 0), (1, -1), SLATE),
        ("FONTNAME", (1, 0), (1, -1), "Helvetica-Bold"),
        ("TOPPADDING", (0, 0), (-1, -1), 6),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
        ("LINEBELOW", (0, 0), (-1, -2), 0.4, colors.HexColor("#e2e8f0")),
    ]))
    elements.append(cover)
    elements.append(Spacer(1, 1.5 * cm))
    elements.append(Paragraph(
        "Confidentiel — diffusion restreinte. Document d'aide à la décision : "
        "les décisions restent prises par les personnes habilitées.", small))
    elements.append(PageBreak())

    # ===== PAGE 2 — Vue d'ensemble =====
    elements.append(Paragraph("1. Vue d'ensemble", h1))
    elements.append(Paragraph("État de la campagne — indicateurs clés", h1sub))
    c = [
        kpi_cell(f"{vol.get('analyzed', 0)} / {vol.get('total', 0)}", "Objectifs analysés", "analyzed",
                 f"{vol.get('analyzed', 0)} objectifs ont été analysés sur {vol.get('total', 0)} importés."),
        kpi_cell(pctv(sr.get("rate")), f"Taux d'objectifs SMART (cible {thr.get('smart', 80)}%)", "smart_rate",
                 f"{sr.get('ok', 0)} objectifs sur {sr.get('total', 0)} atteignent le seuil de conformité SMART.", EMERALD),
        kpi_cell(pctv(avg.get("avg_smart")), "Score SMART moyen", "avg_smart",
                 f"Le score SMART moyen s'établit à {pctv(avg.get('avg_smart'))}."),
        kpi_cell(pctv(avg.get("avg_alignment")), "Alignement stratégique moyen", "avg_alignment",
                 f"L'alignement stratégique moyen s'établit à {pctv(avg.get('avg_alignment'))}."),
        kpi_cell(to_adjust, "Objectifs à ajuster", "to_adjust",
                 f"{to_adjust} objectifs se situent sous le seuil de conformité SMART.", ROSE),
        kpi_cell(lc.get("count", 0), "Analyses à faible confiance IA", "low_confidence",
                 f"{lc.get('count', 0)} analyses présentent une confiance IA inférieure à "
                 f"{thr.get('ai_confidence_min', 0.6)} et appellent une relecture humaine.", VIOLET),
    ]
    grid = Table([[c[0], c[1]], [c[2], c[3]], [c[4], c[5]]], colWidths=[8.4 * cm, 8.4 * cm])
    grid.setStyle(TableStyle([
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 8), ("RIGHTPADDING", (0, 0), (-1, -1), 8),
        ("TOPPADDING", (0, 0), (-1, -1), 10), ("BOTTOMPADDING", (0, 0), (-1, -1), 10),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#e2e8f0")),
    ]))
    elements.append(grid)
    elements.append(PageBreak())

    # ===== PAGE 3 — Qualité SMART =====
    elements.append(Paragraph("2. Qualité SMART", h1))
    elements.append(Paragraph("Conformité des objectifs aux 5 critères", h1sub))
    fails = {}
    for bp in (data.get("blocking_points") or []):
        crit = (bp.get("criteria") or "")
        fails[crit[:1].upper()] = int(bp.get("failures") or 0)
    rows = [["Critère", "Taux de conformité", "Objectifs en échec"]]
    for label, key in [("S — Spécifique", "avg_s"), ("M — Mesurable", "avg_m"), ("A — Atteignable", "avg_a"),
                       ("R — Réaliste", "avg_r"), ("T — Temporel", "avg_t")]:
        rows.append([label, pctv((avg.get(key) or 0) * 100), str(fails.get(label[0], 0))])
    elements.append(table(rows, [7 * cm, 5 * cm, 4.6 * cm]))
    elements.append(Spacer(1, 0.4 * cm))
    sd = data.get("smart_distribution") or []
    if sd:
        elements.append(Paragraph("Répartition des scores SMART", kpilbl))
        elements.append(table([["Niveau", "Objectifs"]] + [[d.get("level", ""), str(d.get("count", 0))] for d in sd],
                              [11.6 * cm, 5 * cm]))
        elements.append(Spacer(1, 0.4 * cm))
    top = max(fails, key=fails.get) if fails else None
    elements.append(para("quality",
                         f"Le critère le plus souvent non satisfait est « {top} » avec {fails.get(top, 0)} objectifs en échec."
                         if top else "Les cinq critères SMART sont évalués sur l'ensemble des objectifs analysés."))
    elements.append(PageBreak())

    # ===== PAGE 4 — Alignement & couverture stratégique =====
    elements.append(Paragraph("3. Alignement &amp; couverture stratégique", h1))
    elements.append(Paragraph("Répartition des objectifs sur les axes stratégiques", h1sub))
    rows = [["Axe stratégique", "Objectifs", "Part (%)", "Signal"]]
    for ax in axes:
        signal = "Orphelin" if ax.get("orphan") else ("Sur-concentré" if ax.get("over_concentrated") else "—")
        rows.append([ax.get("axis_title", ""), str(ax.get("count", 0)), f"{ax.get('share_pct', 0)}%", signal])
    elements.append(table(rows, [8.6 * cm, 2.6 * cm, 2.6 * cm, 2.8 * cm]))
    elements.append(Spacer(1, 0.4 * cm))
    ad = data.get("alignment_distribution") or []
    if ad:
        elements.append(Paragraph("Niveaux d'alignement", kpilbl))
        elements.append(table([["Niveau", "Objectifs"]] + [[d.get("level", ""), str(d.get("count", 0))] for d in ad],
                              [11.6 * cm, 5 * cm]))
        elements.append(Spacer(1, 0.4 * cm))
    elements.append(para("coverage",
                         f"Les objectifs analysés se répartissent sur {len([a for a in axes if a.get('count')])} axe(s) ; "
                         f"le seuil de sur-concentration est fixé à {thr.get('axis_over_concentration', 40)}%."))
    elements.append(PageBreak())

    # ===== PAGE 5 — Points d'attention =====
    elements.append(Paragraph("4. Points d'attention", h1))
    elements.append(Paragraph("Signaux à surveiller — formulés en constats factuels", h1sub))
    elements.append(table([
        ["Signal", "Valeur"],
        ["Objectifs à ajuster", str(to_adjust)],
        ["Analyses à faible confiance IA", str(lc.get("count", 0))],
        ["Axes stratégiques orphelins", str(orphans)],
        ["Axes sur-concentrés", str(over)],
    ], [11.6 * cm, 5 * cm]))
    elements.append(Spacer(1, 0.4 * cm))
    elements.append(para("attention",
                         f"{to_adjust} objectifs sont à ajuster, {lc.get('count', 0)} analyses appellent une "
                         f"relecture humaine, et {orphans} axe(s) stratégique(s) ne reçoivent aucun objectif."))
    elements.append(PageBreak())

    # ===== PAGE 6 — Synthèse générale + annexes =====
    elements.append(Paragraph("5. Synthèse générale", h1))
    fallback_syn = (
        f"Sur {vol.get('total', 0)} objectifs importés, {vol.get('analyzed', 0)} ont été analysés. "
        f"Le taux d'objectifs SMART atteint {pctv(sr.get('rate'))} (cible {thr.get('smart', 80)}%) "
        f"et l'alignement moyen {pctv(avg.get('avg_alignment'))}. "
        f"{to_adjust} objectifs sont à ajuster et {lc.get('count', 0)} analyses appellent une relecture humaine.")
    elements.append(Paragraph(_esc(synthesis or fallback_syn), body))
    elements.append(Spacer(1, 0.7 * cm))
    elements.append(Paragraph("Annexes — disponibles sur demande auprès de la RH", h1sub))
    for a in [
        "Liste détaillée des objectifs à ajuster (sans KPI, sans échéance, flous, désalignés).",
        "Propositions de reformulation IA (avant / après), soumises à validation humaine.",
        "Méthodologie de calcul des scores SMART et d'alignement, et seuils appliqués.",
    ]:
        elements.append(Paragraph(f"• {a}", small))
        elements.append(Spacer(1, 0.15 * cm))
    elements.append(Spacer(1, 0.7 * cm))
    elements.append(Paragraph(
        "Confidentiel — diffusion restreinte · Éléments fournis à titre d'aide à la décision.", small))

    doc.build(elements)
    output.seek(0)
    filename = f"Rapport_COMEX_BMCI_{datetime.now().strftime('%Y%m%d_%H%M%S')}.pdf"
    return StreamingResponse(
        output,
        media_type="application/pdf",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )
