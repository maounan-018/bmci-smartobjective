#!/usr/bin/env python3
"""
services/hr_import_service.py — Imports Excel du domaine RH.

RÔLE
    Import des objectifs stratégiques d'entreprise et des objectifs
    collaborateurs (avec création des comptes manquants). ISOLATION STRICTE :
    périmètre HR uniquement.

EMPLACEMENT MVC
    Couche "services" : logique métier RH (imports).
"""

import io
import logging
import secrets
from typing import List

from fastapi import HTTPException

from app.core.security import hash_password
from app.database.session import db_session
from app.repositories import (
    user_repository,
    objective_repository,
    company_objective_repository,
    hr_import_repository,
)
from app.services.hr_common import require_hr

logger = logging.getLogger(__name__)

try:
    import openpyxl
    EXCEL_SUPPORT = True
except ImportError:  # pragma: no cover
    EXCEL_SUPPORT = False


def import_strategy(filename: str, content: bytes, current_user: dict) -> dict:
    """Import Excel des objectifs stratégiques d'entreprise."""
    require_hr(current_user)
    if not EXCEL_SUPPORT:
        raise HTTPException(status_code=500, detail="openpyxl non installé. Exécutez: pip install openpyxl")

    try:
        wb = openpyxl.load_workbook(io.BytesIO(content))
        ws = wb.active
        headers = [str(cell.value or "").lower().strip() for cell in ws[1]]

        col_map = {}
        for idx, h in enumerate(headers):
            if h in ["id", "reference", "ref", "code"]:
                col_map["id"] = idx
            elif h in ["titre", "title", "objectif", "axe", "axe stratégique", "axe strategique"]:
                col_map["title"] = idx
            elif h in ["description", "desc", "detail", "détail"]:
                col_map["description"] = idx
            elif h in ["departement", "department", "direction", "entité", "entite", "entity"]:
                col_map["department"] = idx
            elif h in ["priorité", "priorite", "priority", "poids", "weight"]:
                col_map["priority"] = idx

        if "title" not in col_map:
            raise HTTPException(status_code=400, detail="Colonne 'titre' ou 'title' obligatoire dans le fichier Excel.")

        batch_id = f"strat-{secrets.token_hex(4)}"
        imported = 0

        with db_session() as db:
            for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
                title = row[col_map["title"]] if col_map.get("title") is not None else None
                if not title or str(title).strip() == "":
                    continue

                obj_id = row[col_map["id"]] if col_map.get("id") is not None else f"STRAT-{row_idx}"
                description = row[col_map["description"]] if col_map.get("description") is not None else ""
                department = row[col_map["department"]] if col_map.get("department") is not None else ""
                priority = row[col_map["priority"]] if col_map.get("priority") is not None else 1
                try:
                    priority = int(priority) if priority else 1
                except (ValueError, TypeError):
                    priority = 1

                company_objective_repository.merge(db, {
                    "id": str(obj_id).strip(),
                    "title": str(title).strip(),
                    "description": str(description or "").strip(),
                    "department": str(department or "").strip(),
                    "priority": priority,
                    "created_by": current_user["id"],
                    "import_batch_id": batch_id,
                })
                imported += 1

            hr_import_repository.create(
                db, import_id=batch_id, import_type="strategy",
                filename=filename, records_count=imported, user_id=current_user["id"],
            )

        return {"success": True, "batch_id": batch_id, "imported": imported, "filename": filename}

    except HTTPException:
        raise
    except Exception as e:
        logger.error("Erreur import RH: %s", e)
        raise HTTPException(status_code=400, detail="Erreur lors de l'import du fichier. Vérifiez le format et réessayez.")


def import_objectives(filename: str, content: bytes, current_user: dict) -> dict:
    """
    Import Excel des objectifs collaborateurs.

    SÉCURITÉ (F10) : les comptes créés au passage reçoivent un mot de passe
    aléatoire FORT par utilisateur (token_urlsafe), jamais un seed partagé. Les
    credentials temporaires sont renvoyés dans `created_users` pour transmission
    individuelle (à changer à la première connexion).
    """
    require_hr(current_user)
    if not EXCEL_SUPPORT:
        raise HTTPException(status_code=500, detail="openpyxl non installé. Exécutez: pip install openpyxl")

    try:
        wb = openpyxl.load_workbook(io.BytesIO(content))
        ws = wb.active
        headers = [str(cell.value or "").lower().strip() for cell in ws[1]]

        col_map = {}
        for idx, h in enumerate(headers):
            if h in ["id", "reference", "ref"]:
                col_map["id"] = idx
            elif h in ["collaborateur", "collab", "user", "utilisateur", "nom", "name", "employee"]:
                col_map["user_name"] = idx
            elif h in ["email", "mail", "e-mail"]:
                col_map["email"] = idx
            elif h in ["titre", "title", "objectif", "objective"]:
                col_map["title"] = idx
            elif h in ["description", "desc", "detail", "détail"]:
                col_map["description"] = idx
            elif h in ["entité", "entite", "entity", "direction", "department", "service"]:
                col_map["entity"] = idx
            elif h in ["manager", "responsable", "n+1", "superviseur"]:
                col_map["manager"] = idx
            elif h in ["période", "periode", "period", "année", "annee", "year", "trimestre", "quarter"]:
                col_map["period"] = idx
            elif h in ["contexte", "context"]:
                col_map["context"] = idx

        if "title" not in col_map:
            raise HTTPException(status_code=400, detail="Colonne 'titre' ou 'title' obligatoire dans le fichier Excel.")

        batch_id = f"obj-{secrets.token_hex(4)}"
        imported = 0
        users_created = 0
        created_users: List[dict] = []

        with db_session() as db:
            for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
                title = row[col_map["title"]] if col_map.get("title") is not None else None
                if not title or str(title).strip() == "":
                    continue

                user_name = row[col_map["user_name"]] if col_map.get("user_name") is not None else f"Collaborateur-{row_idx}"
                email = row[col_map["email"]] if col_map.get("email") is not None else None
                entity = row[col_map["entity"]] if col_map.get("entity") is not None else ""
                manager = row[col_map["manager"]] if col_map.get("manager") is not None else ""
                period = row[col_map["period"]] if col_map.get("period") is not None else ""
                description = row[col_map["description"]] if col_map.get("description") is not None else ""
                context = row[col_map["context"]] if col_map.get("context") is not None else ""
                obj_id_raw = row[col_map["id"]] if col_map.get("id") is not None else None

                if email:
                    existing = user_repository.get_by_email(db, str(email).strip())
                    if existing:
                        user_id = existing.id
                    else:
                        user_id = f"user-{secrets.token_hex(4)}"
                        temp_password = secrets.token_urlsafe(16)
                        user_repository.create(
                            db, user_id=user_id, email=str(email).strip(),
                            password_hash=hash_password(temp_password),
                            name=str(user_name).strip(), role="COLLAB",
                            entity=str(entity or "").strip(),
                        )
                        users_created += 1
                        created_users.append({
                            "user_id": user_id,
                            "name": str(user_name).strip(),
                            "email": str(email).strip(),
                            "temporary_password": temp_password,
                        })
                else:
                    user_id = f"import-user-{row_idx}"
                    if not user_repository.get_by_id(db, user_id):
                        temp_password = secrets.token_urlsafe(16)
                        temp_email = f"import_{row_idx}_{batch_id}@temp.local"
                        user_repository.create(
                            db, user_id=user_id, email=temp_email,
                            password_hash=hash_password(temp_password),
                            name=str(user_name).strip(), role="COLLAB",
                            entity=str(entity or "").strip(),
                        )
                        users_created += 1
                        created_users.append({
                            "user_id": user_id,
                            "name": str(user_name).strip(),
                            "email": temp_email,
                            "temporary_password": temp_password,
                            "note": "Email synthetique : a editer avant transmission.",
                        })

                obj_id = str(obj_id_raw).strip() if obj_id_raw else f"obj-{secrets.token_hex(6)}"
                objective_repository.merge(db, {
                    "id": obj_id,
                    "user_id": user_id,
                    "title": str(title).strip(),
                    "description": str(description or "").strip(),
                    "context": str(context or "").strip(),
                    "entity": str(entity or "").strip(),
                    "manager_name": str(manager or "").strip(),
                    "period": str(period or "").strip(),
                    "import_batch_id": batch_id,
                    "data_scope": "HR",
                    "status": "draft",
                    "progress": "not_started",
                })
                imported += 1

            hr_import_repository.create(
                db, import_id=batch_id, import_type="objectives",
                filename=filename, records_count=imported, user_id=current_user["id"],
            )

        return {
            "success": True,
            "batch_id": batch_id,
            "imported": imported,
            "users_created": users_created,
            "filename": filename,
            "created_users": created_users,
        }

    except HTTPException:
        raise
    except Exception as e:
        logger.error("Erreur import RH: %s", e)
        raise HTTPException(status_code=400, detail="Erreur lors de l'import du fichier. Vérifiez le format et réessayez.")
